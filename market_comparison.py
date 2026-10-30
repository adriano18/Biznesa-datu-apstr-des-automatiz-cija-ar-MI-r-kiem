"""Conservative comparison of identical items in two shop samples.

A price comparison is produced only for the same currency, equivalent
packaging and a strong product identity match (GTIN or exact normalized name).
No fuzzy similarity guesses are used.
"""
from datetime import datetime
from pathlib import Path
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent
REPORT_DIR = ROOT / "reports"

MATCH_COLUMNS = [
    "product", "shop_a", "shop_b", "price_a", "price_b", "currency",
    "weight_kg", "difference_eur", "difference_pct_vs_b", "cheaper",
    "availability_a", "availability_b", "url_a", "url_b", "match_method"
]


def _clean(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def _plain(value):
    text = unicodedata.normalize("NFKD", _clean(value).casefold())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"(?<!\w)\d+(?:[.,]\d+)?\s*(?:kg|g)\b", " ", text)
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _weight(row):
    try:
        value = float(row.get("weight_kg", float("nan")))
        return value if pd.notna(value) and value > 0 else None
    except (TypeError, ValueError):
        return None


def _same_package(a, b):
    wa, wb = _weight(a), _weight(b)
    return (wa is None and wb is None) or (
        wa is not None and wb is not None and abs(wa - wb) < 0.001
    )


def _match_method(a, b):
    if not _same_package(a, b):
        return None
    ga, gb = _clean(a.get("gtin")), _clean(b.get("gtin"))
    if ga and gb:
        # Different GTINs mean different sellable units / variants.
        return "GTIN" if ga == gb else None
    na, nb = _plain(a.get("name")), _plain(b.get("name"))
    if na and na == nb:
        return "Exact normalized name"
    return None


def match_products(df_a, df_b):
    """Return price comparisons only for strongly matched products.

    Difference = price B - price A, relative percentage uses B as base.
    Positive difference means shop A has the lower listed price.
    """
    rows = []
    used_b = set()

    for _, a in df_a.iterrows():

        for j, b in df_b.iterrows():
            if j in used_b:
                continue
            if _clean(a.get("currency")) != _clean(b.get("currency")):
                continue
            if _clean(a.get("currency")) != "EUR":
                continue

            method = _match_method(a, b)

            if method is None:
                continue

            pa = pd.to_numeric(a.get("price"), errors="coerce")
            pb = pd.to_numeric(b.get("price"), errors="coerce")

            if pd.isna(pa) or pd.isna(pb) or pa < 0 or pb <= 0:
                continue

            difference = round(float(pb) - float(pa), 2)
            percent = round(100 * difference / float(pb), 2)
            cheaper = "Equal price" if difference == 0 else (
                "Shop A" if difference > 0 else "Shop B"
            )

            rows.append({
                "product": a.get("name", ""),
                "shop_a": a.get("store", "Shop A"),
                "shop_b": b.get("store", "Shop B"),
                "price_a": round(float(pa), 2),
                "price_b": round(float(pb), 2),
                "currency": "EUR",
                "weight_kg": _weight(a),
                "difference_eur": difference,
                "difference_pct_vs_b": percent,
                "cheaper": cheaper,
                "availability_a": a.get("availability", ""),
                "availability_b": b.get("availability", ""),
                "url_a": a.get("url", ""),
                "url_b": b.get("url", ""),
                "match_method": method,
            })

            used_b.add(j)

            break

    return pd.DataFrame(rows, columns=MATCH_COLUMNS)


def save_market_excel(df_a, df_b, matches):
    """Generate an audit-friendly Excel with all collected products and matches."""
    REPORT_DIR.mkdir(exist_ok=True)

    path = REPORT_DIR / f"market_comparison_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Metric", "Shop A", "Shop B"])

    for label, a, b in [
        ("Store", df_a["store"].iloc[0], df_b["store"].iloc[0]),
        ("Products in sample", len(df_a), len(df_b)),
        ("Average price EUR", float(df_a["price"].mean()), float(df_b["price"].mean())),
        ("Median EUR/kg", float(df_a["price_per_kg"].median()), float(df_b["price_per_kg"].median())),
        ("Matched identical products", len(matches), len(matches)),
    ]:
        summary.append([label, a, b])
    summary.append(["Important", "Different products cannot be directly compared by average price."])

    all_products = pd.concat([df_a, df_b], ignore_index=True)

    # Brand names are taken from Product JSON-LD where present, not guessed.
    if "brand" in all_products.columns:
        brand_data = all_products.copy()
        brand_data["brand"] = brand_data["brand"].fillna("Unknown")
        brand_counts = brand_data.groupby(["store", "brand"], dropna=False).agg(
            products=("name", "size"), average_price_eur=("price", "mean")
        ).reset_index()
    else:
        brand_counts = pd.DataFrame(columns=["store", "brand", "products", "average_price_eur"])

    for title, data in (("Products", all_products),
                        ("Brands", brand_counts),
                        ("Exact Matches", matches)):
        ws = wb.create_sheet(title)
        ws.append(list(data.columns))

        for values in data.itertuples(index=False, name=None):
            ws.append([None if pd.isna(v) else v for v in values])

        if title == "Exact Matches" and matches.empty:
            ws.append(["No confidently matched identical products in this sample"])
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{ws.cell(1,ws.max_column).column_letter}{ws.max_row}"

        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = min(65, max(16, max(len(str(c.value or "")) for c in col[:20]) + 2))

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"  # avoid spreadsheet formula injection
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243247")
        ws.freeze_panes = "B2"

    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 42
    summary.column_dimensions["C"].width = 42

    wb.save(path)

    return str(path)


if __name__ == "__main__":
    # Smoke test against two local CSV exports.
    paths = sorted(REPORT_DIR.glob("products_*.csv"))

    if len(paths) < 2:
        raise SystemExit("Two product CSV files are required in reports/.")
    
    a = pd.read_csv(paths[0], dtype={"sku": "string", "gtin": "string"})
    b = pd.read_csv(paths[1], dtype={"sku": "string", "gtin": "string"})
    
    matched = match_products(a, b)

    print("Confirmed matches:", len(matched))
    print("Excel:", save_market_excel(a, b, matched))
