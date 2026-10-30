"""Export the two-shop SEO audit to a formatted Excel workbook.

Uses the existing seo_analyzer results and stores status/value pairs for
both shops. REVIEW and INFO are diagnostic labels, not Google rankings.
The report is separate from the product comparison Excel."""


from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from seo_analyzer import analyze_seo


# STEP 1: Turn each fetched HTML response into a structured SEO audit.
def save_excel(url_a, data_a, url_b, data_b):

    def get_audit(data):
        if data["status"] == "OK":
            return analyze_seo(data["html"])

        return {
            "Fetch": {
                "status": data["status"],
                "value": data["message"]
            }
        }

    # Both shops are represented even if one page could not be downloaded.
    audit_a = get_audit(data_a)
    audit_b = get_audit(data_b)

    # STEP 2: Build one side-by-side worksheet for SEO parameters.
    wb = Workbook()
    ws = wb.active
    ws.title = "SEO Comparison"

    ws.append(["Shop A URL", url_a])
    ws.append(["Shop B URL", url_b])
    ws.append([])
    ws.append(["Parameter", "A Status", "A Value", "B Status", "B Value"])

    parameters = list(dict.fromkeys([*audit_a, *audit_b]))

    for parameter in parameters:
        a = audit_a.get(parameter, {"status": "NOT CHECKED", "value": ""})
        b = audit_b.get(parameter, {"status": "NOT CHECKED", "value": ""})

        ws.append([
            parameter,
            a["status"],
            a["value"],
            b["status"],
            b["value"]
        ])

    # STEP 3: Make each audit status visually distinguishable in Excel.
    colors = {
        "OK": "DDF2E2",
        "REVIEW": "FFF0C2",
        "ERROR": "FADBD8",
        "INFO": "E3EDF9",
        "NOT CHECKED": "EAEAEA",
        "SKIPPED": "EAEAEA"
    }

    for cell in ws[4]:
        cell.fill = PatternFill("solid", fgColor="243247")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)

    for row in range(5, ws.max_row + 1):
        for column in (2, 4):
            cell = ws.cell(row, column)
            color = colors.get(cell.value, "FFFFFF")
            cell.fill = PatternFill("solid", fgColor=color)

        for column in (3, 5):
            cell = ws.cell(row, column)
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for column, width in {
        "A": 28,
        "B": 17,
        "C": 60,
        "D": 17,
        "E": 60
    }.items():
        ws.column_dimensions[column].width = width

    # Keep column labels visible and enable sorting/filtering.
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:E{ws.max_row}"

    # STEP 4: Export a timestamped workbook to the reports directory.
    folder = Path("reports")
    folder.mkdir(exist_ok=True)

    filename = folder / f"seo_comparison_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

    wb.save(filename)

    return str(filename.resolve())
